from flask import Flask, render_template, request, jsonify, send_file
import os
import datetime
import sqlite3
import json
import threading
import io
import re

app = Flask(__name__)

# === 選用依賴 ===
try:
    from flask_sock import Sock
    sock = Sock(app)
    WEBSOCKET_OK = True
except ImportError:
    WEBSOCKET_OK = False
    sock = None

try:
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler
    WATCHDOG_OK = True
except ImportError:
    WATCHDOG_OK = False

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

DEFAULT_ROOT = os.path.expanduser("~")
INDEX_DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'search_index.db')

# WebSocket 連線池（執行緒安全）
_ws_clients = set()
_ws_lock = threading.Lock()

# 檔案監控
_observer = None
_observer_root = None
_observer_lock = threading.Lock()


def _broadcast(event_type, path):
    """廣播檔案事件到所有 WebSocket 連線"""
    msg = json.dumps({'event': event_type, 'path': path.replace('\\', '/')})
    with _ws_lock:
        dead = []
        for ws in list(_ws_clients):
            try:
                ws.send(msg)
            except Exception:
                dead.append(ws)
        for ws in dead:
            _ws_clients.discard(ws)


if WATCHDOG_OK:
    class _MDHandler(FileSystemEventHandler):
        def on_modified(self, event):
            if not event.is_directory and str(event.src_path).lower().endswith('.md'):
                _broadcast('changed', event.src_path)

        def on_created(self, event):
            if not event.is_directory and str(event.src_path).lower().endswith('.md'):
                _broadcast('created', event.src_path)

        def on_deleted(self, event):
            if not event.is_directory and str(event.src_path).lower().endswith('.md'):
                _broadcast('deleted', event.src_path)

        def on_moved(self, event):
            if not event.is_directory:
                _broadcast('moved', event.dest_path)


def start_watcher(root):
    """啟動目錄監控"""
    global _observer, _observer_root
    if not WATCHDOG_OK:
        return
    with _observer_lock:
        if _observer and _observer.is_alive():
            _observer.stop()
            _observer.join(timeout=2)
        _observer = Observer()
        _observer.schedule(_MDHandler(), root, recursive=True)
        _observer.daemon = True
        _observer.start()
        _observer_root = root


# === 工具函數 ===
def safe_path(root, path):
    abs_root = os.path.abspath(root)
    abs_path = os.path.abspath(path)
    return abs_path if abs_path.startswith(abs_root) else None


def scan_tree(directory):
    tree = []
    try:
        entries = sorted(os.listdir(directory),
                         key=lambda x: (not os.path.isdir(os.path.join(directory, x)), x.lower()))
    except PermissionError:
        return tree
    for entry in entries:
        full_path = os.path.join(directory, entry)
        if os.path.isdir(full_path):
            if entry.startswith('.') or entry in ('node_modules', '__pycache__', '.git', '.venv', 'venv'):
                continue
            children = scan_tree(full_path)
            if children:
                tree.append({'name': entry, 'type': 'directory',
                             'path': full_path.replace('\\', '/'), 'children': children})
        elif entry.lower().endswith('.md'):
            tree.append({'name': entry, 'type': 'file', 'path': full_path.replace('\\', '/')})
    return tree


# === SQLite FTS5 全文索引 ===
def _init_index_db():
    conn = sqlite3.connect(INDEX_DB)
    conn.execute('''CREATE VIRTUAL TABLE IF NOT EXISTS docs USING fts5(
        path UNINDEXED, name UNINDEXED, content,
        tokenize="unicode61 remove_diacritics 1"
    )''')
    conn.execute('CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT)')
    conn.commit()
    return conn


def rebuild_index(root):
    conn = _init_index_db()
    conn.execute('DELETE FROM docs')
    conn.execute('INSERT OR REPLACE INTO meta VALUES ("root", ?)', (os.path.abspath(root).replace('\\', '/'),))
    count = 0
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith('.')
                       and d not in ('node_modules', '__pycache__', '.git', '.venv', 'venv')]
        for fname in filenames:
            if fname.lower().endswith('.md'):
                fpath = os.path.join(dirpath, fname)
                try:
                    with open(fpath, 'r', encoding='utf-8') as f:
                        content = f.read()
                    conn.execute('INSERT INTO docs VALUES (?, ?, ?)',
                                 (fpath.replace('\\', '/'), fname, content))
                    count += 1
                except Exception:
                    pass
    conn.commit()
    conn.close()
    return count


def search_fts5(root, query, max_files=50, max_per_file=10):
    """使用 FTS5 索引搜尋。若索引不可用回傳 None"""
    try:
        conn = _init_index_db()
        row = conn.execute('SELECT value FROM meta WHERE key="root"').fetchone()
        if not row:
            conn.close()
            return None
        indexed_root = row[0]
        abs_root = os.path.abspath(root).replace('\\', '/')
        if not abs_root.startswith(indexed_root) and indexed_root != abs_root:
            conn.close()
            return None
        root_prefix = abs_root if abs_root.endswith('/') else abs_root + '/'
        safe_query = query.replace('"', '""')
        rows = conn.execute(
            'SELECT path, name, snippet(docs, 2, "[[", "]]", "…", 24) '
            'FROM docs WHERE content MATCH ? AND (path = ? OR path LIKE ?) LIMIT ?',
            (f'"{safe_query}"', abs_root, root_prefix + '%', max_files * max_per_file)
        ).fetchall()
        conn.close()
        seen = {}
        for path, name, snippet in rows:
            if path not in seen:
                seen[path] = {'path': path, 'name': name, 'matches': []}
            if len(seen[path]['matches']) < max_per_file:
                seen[path]['matches'].append({'line': 0, 'text': snippet})
        return list(seen.values())[:max_files]
    except Exception:
        return None


def get_index_info():
    try:
        conn = _init_index_db()
        row = conn.execute('SELECT value FROM meta WHERE key="root"').fetchone()
        count = conn.execute('SELECT COUNT(*) FROM docs').fetchone()[0]
        conn.close()
        return {'root': row[0] if row else None, 'count': count}
    except Exception:
        return {'root': None, 'count': 0}


# === Markdown → DOCX 轉換 ===
def _add_inline(paragraph, text):
    pattern = r'(\*\*[^*\n]+\*\*|\*[^*\n]+\*|`[^`\n]+`|\[[^\]\n]+\]\([^)\n]+\)|[^*`\[\n]+)'
    for chunk in re.findall(pattern, text):
        if chunk.startswith('**') and chunk.endswith('**') and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        elif chunk.startswith('`') and chunk.endswith('`') and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        elif chunk.startswith('['):
            link_text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', chunk)
            run = paragraph.add_run(link_text)
            if DOCX_OK:
                run.font.color.rgb = RGBColor(9, 105, 218)
        else:
            stripped = chunk.strip()
            if stripped:
                paragraph.add_run(chunk)


def md_to_docx(content):
    if not DOCX_OK:
        return None
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    lines = content.split('\n')
    i = 0
    in_code = False
    code_lines = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]

        # 程式碼區塊
        if line.startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                code_text = '\n'.join(code_lines) or ' '
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'F6F8FA')
                pPr.append(shd)
                ind = OxmlElement('w:ind')
                ind.set(qn('w:left'), '360')
                pPr.append(ind)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # 表格
        if '|' in line and line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            if re.match(r'^\s*\|[\s|:=-]+\|\s*$', line):
                i += 1
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            table_rows.append(cells)
            i += 1
            continue
        else:
            if in_table and table_rows:
                max_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=max_cols)
                tbl.style = 'Table Grid'
                for r_idx, row_cells in enumerate(table_rows):
                    for c_idx, cell_text in enumerate(row_cells[:max_cols]):
                        cell = tbl.cell(r_idx, c_idx)
                        cell.text = cell_text
                        if r_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
                in_table = False
                table_rows = []

        # 標題
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            text = re.sub(r'[*`]', '', m.group(2))
            doc.add_heading(text, level=level)
            i += 1
            continue

        # 水平線
        if re.match(r'^[-*_]{3,}\s*$', line):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            btm = OxmlElement('w:bottom')
            btm.set(qn('w:val'), 'single')
            btm.set(qn('w:sz'), '4')
            btm.set(qn('w:color'), 'D0D7DE')
            pBdr.append(btm)
            pPr.append(pBdr)
            i += 1
            continue

        # 引用
        if line.startswith('>'):
            text = line[1:].strip()
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            ind = OxmlElement('w:ind')
            ind.set(qn('w:left'), '720')
            pPr.append(ind)
            pBdr = OxmlElement('w:pBdr')
            left_bdr = OxmlElement('w:left')
            left_bdr.set(qn('w:val'), 'single')
            left_bdr.set(qn('w:sz'), '12')
            left_bdr.set(qn('w:color'), '0969DA')
            pBdr.append(left_bdr)
            pPr.append(pBdr)
            _add_inline(p, text)
            i += 1
            continue

        # 無序清單
        m = re.match(r'^\s*[*\-+]\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            _add_inline(p, m.group(1))
            i += 1
            continue

        # 有序清單
        m = re.match(r'^\s*\d+\.\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            _add_inline(p, m.group(1))
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 一般段落
        p = doc.add_paragraph()
        _add_inline(p, line)
        i += 1

    # 清除剩餘表格
    if in_table and table_rows:
        max_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=max_cols)
        tbl.style = 'Table Grid'
        for r_idx, row_cells in enumerate(table_rows):
            for c_idx, cell_text in enumerate(row_cells[:max_cols]):
                tbl.cell(r_idx, c_idx).text = cell_text

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# === Flask 路由 ===
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/tree')
def api_tree():
    root = request.args.get('root', DEFAULT_ROOT)
    if not os.path.isdir(root):
        return jsonify({'error': '目錄不存在'}), 400
    start_watcher(root)
    tree = scan_tree(root)
    return jsonify({'root': root.replace('\\', '/'), 'tree': tree,
                    'websocket': WEBSOCKET_OK, 'watchdog': WATCHDOG_OK})


@app.route('/api/file', methods=['GET'])
def api_file():
    path = request.args.get('path', '')
    if not path or not os.path.isfile(path):
        return jsonify({'error': '檔案不存在'}), 400
    try:
        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    stat = os.stat(path)
    mod_time = datetime.datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M:%S')
    lines = content.count('\n') + (1 if content and not content.endswith('\n') else 0)
    word_count = len(content) - content.count('\n') - content.count('\r') - content.count(' ')
    return jsonify({'content': content, 'wordCount': word_count,
                    'lineCount': lines, 'modifiedTime': mod_time,
                    'fileName': os.path.basename(path)})


@app.route('/api/file/save', methods=['POST'])
def api_save_file():
    """儲存檔案內容"""
    data = request.get_json()
    if not data:
        return jsonify({'error': '無效請求'}), 400
    path = data.get('path', '').strip()
    content = data.get('content', '')
    if not path:
        return jsonify({'error': '路徑不可為空'}), 400
    if not os.path.isfile(path):
        return jsonify({'error': '檔案不存在'}), 400
    try:
        with open(path, 'w', encoding='utf-8') as f:
            f.write(content)
        stat = os.stat(path)
        mod_time = datetime.datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M:%S')
        lines = content.count('\n') + (1 if content and not content.endswith('\n') else 0)
        word_count = len(content) - content.count('\n') - content.count('\r') - content.count(' ')
        return jsonify({'success': True, 'modifiedTime': mod_time,
                        'wordCount': word_count, 'lineCount': lines})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/search')
def api_search():
    root = request.args.get('root', DEFAULT_ROOT)
    query = request.args.get('q', '').strip()
    use_index = request.args.get('index', 'true').lower() == 'true'
    if not query:
        return jsonify({'results': [], 'indexed': False})
    if use_index:
        results = search_fts5(root, query)
        if results is not None:
            return jsonify({'results': results, 'indexed': True})
    # 線性搜尋後備
    results = []
    query_lower = query.lower()

    def search_dir(directory):
        try:
            entries = os.listdir(directory)
        except PermissionError:
            return
        for entry in entries:
            full_path = os.path.join(directory, entry)
            if os.path.isdir(full_path):
                if entry.startswith('.') or entry in ('node_modules', '__pycache__', '.git', '.venv', 'venv'):
                    continue
                search_dir(full_path)
            elif entry.lower().endswith('.md'):
                try:
                    with open(full_path, 'r', encoding='utf-8') as f:
                        file_lines = f.readlines()
                    matches = []
                    for ln, line in enumerate(file_lines, 1):
                        if query_lower in line.lower():
                            matches.append({'line': ln, 'text': line.strip()[:200]})
                    if matches:
                        results.append({'path': full_path.replace('\\', '/'),
                                        'name': entry, 'matches': matches[:10]})
                except Exception:
                    pass

    search_dir(root)
    return jsonify({'results': results[:50], 'indexed': False})


@app.route('/api/index/rebuild', methods=['POST'])
def api_rebuild_index():
    data = request.get_json() or {}
    root = data.get('root', DEFAULT_ROOT)
    if not os.path.isdir(root):
        return jsonify({'error': '目錄不存在'}), 400
    try:
        count = rebuild_index(root)
        return jsonify({'success': True, 'count': count})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/index/info')
def api_index_info():
    return jsonify(get_index_info())


@app.route('/api/export/html')
def api_export_html():
    """回傳 Markdown 原始內容供前端轉成 HTML"""
    path = request.args.get('path', '')
    if not path or not os.path.isfile(path):
        return jsonify({'error': '檔案不存在'}), 400
    try:
        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()
        name = os.path.basename(path).replace('.md', '')
        return jsonify({'content': content, 'name': name})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/export/docx')
def api_export_docx():
    """匯出 Markdown 為 DOCX 格式"""
    if not DOCX_OK:
        return jsonify({'error': 'python-docx 未安裝，請執行: pip install python-docx'}), 500
    path = request.args.get('path', '')
    if not path or not os.path.isfile(path):
        return jsonify({'error': '檔案不存在'}), 400
    try:
        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()
        buf = md_to_docx(content)
        name = os.path.basename(path).replace('.md', '') + '.docx'
        return send_file(buf,
                         mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                         as_attachment=True,
                         download_name=name)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# === WebSocket ===
if WEBSOCKET_OK:
    @sock.route('/ws')
    def websocket(ws):
        with _ws_lock:
            _ws_clients.add(ws)
        try:
            while True:
                try:
                    msg = ws.receive(timeout=30)
                    if msg is None:
                        break
                    if msg == 'ping':
                        ws.send('pong')
                except Exception:
                    break
        finally:
            with _ws_lock:
                _ws_clients.discard(ws)


if __name__ == '__main__':
    app.run(debug=False, port=12017, threaded=True)
